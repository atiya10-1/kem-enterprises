from dotenv import load_dotenv
from pathlib import Path
ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / ".env")

import os
import io
import re
import csv
import uuid
import logging
from datetime import datetime, timezone, timedelta
from typing import List, Optional

from fastapi import FastAPI, APIRouter, HTTPException, Depends, Request, Response, UploadFile, File, Header, Query
from starlette.middleware.cors import CORSMiddleware
from starlette.responses import Response as StarletteResponse
from pydantic import BaseModel, Field

from db import db, client
from auth import (hash_password, verify_password, create_access_token, get_current_user,
                  require_section, seed_admin, user_can, ROLE_PERMISSIONS)
from storage import init_storage, put_object, get_object, APP_NAME, MIME_TYPES
from seed_data import seed_data, slugify, now_iso
import pdfgen

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger("kem")

app = FastAPI(title="KEM Enterprises API")
api = APIRouter(prefix="/api")


# ----------------------------- helpers -----------------------------
def clean(doc):
    if doc:
        doc.pop("_id", None)
    return doc


async def unique_slug(collection, base, exclude_id=None):
    slug = base or "item"
    i = 1
    while True:
        q = {"slug": slug}
        if exclude_id:
            q["id"] = {"$ne": exclude_id}
        if not await collection.find_one(q):
            return slug
        i += 1
        slug = f"{base}-{i}"


async def get_settings():
    s = await db.settings.find_one({"id": "global"})
    return clean(s) if s else {}


async def log_event(event, **data):
    await db.analytics.insert_one({"id": str(uuid.uuid4()), "event": event,
                                   "date": now_iso(), "month": datetime.now(timezone.utc).strftime("%Y-%m"), **data})


async def next_number(kind):
    s = await get_settings()
    cfg = s.get(kind, {})
    prefix = cfg.get("prefix", kind[:3].upper())
    start = cfg.get("start", 1001)
    count = await (db.invoices if kind == "invoice" else db.quotations).count_documents({})
    return f"{prefix}-{start + count}"


# ----------------------------- models -----------------------------
class LoginIn(BaseModel):
    email: str
    password: str


class UserIn(BaseModel):
    name: str
    email: str
    password: str
    role: str = "sales"


class CategoryIn(BaseModel):
    name: str
    description: Optional[str] = ""
    image: Optional[str] = ""
    parent_id: Optional[str] = None
    order: Optional[int] = 0
    active: Optional[bool] = True
    seo_title: Optional[str] = ""
    seo_description: Optional[str] = ""
    seo_keywords: Optional[str] = ""


class ProductIn(BaseModel):
    name: str
    sku: Optional[str] = ""
    category_id: str
    subcategory: Optional[str] = ""
    images: Optional[List[str]] = []
    video: Optional[str] = ""
    short_description: Optional[str] = ""
    description: Optional[str] = ""
    specifications: Optional[List[dict]] = []
    material: Optional[str] = ""
    finish: Optional[str] = ""
    size: Optional[str] = ""
    dimensions: Optional[str] = ""
    wire_diameter: Optional[str] = ""
    load_capacity: Optional[str] = ""
    thread_size: Optional[str] = ""
    application: Optional[str] = ""
    variants: Optional[List[dict]] = []
    packaging: Optional[str] = ""
    moq: Optional[str] = ""
    brand: Optional[str] = "KEM Enterprises"
    tags: Optional[List[str]] = []
    related_ids: Optional[List[str]] = []
    downloads: Optional[List[dict]] = []
    price: Optional[float] = 0
    show_price: Optional[bool] = False
    featured: Optional[bool] = False
    is_new: Optional[bool] = False
    best_seller: Optional[bool] = False
    popular: Optional[bool] = False
    active: Optional[bool] = True
    seo_title: Optional[str] = ""
    seo_description: Optional[str] = ""
    seo_keywords: Optional[str] = ""


class EnquiryIn(BaseModel):
    name: str
    company: Optional[str] = ""
    phone: str
    email: Optional[str] = ""
    whatsapp: Optional[str] = ""
    product_id: Optional[str] = ""
    product_name: Optional[str] = ""
    quantity: Optional[str] = ""
    requirement: Optional[str] = ""
    message: Optional[str] = ""


class EnquiryUpdate(BaseModel):
    status: Optional[str] = None
    assigned_to: Optional[str] = None
    note: Optional[str] = None


class CustomerIn(BaseModel):
    name: str
    company: Optional[str] = ""
    phone: Optional[str] = ""
    whatsapp: Optional[str] = ""
    email: Optional[str] = ""
    address: Optional[str] = ""
    gstin: Optional[str] = ""
    notes: Optional[str] = ""


class QuotationIn(BaseModel):
    customer_id: Optional[str] = ""
    customer_name: str
    company: Optional[str] = ""
    address: Optional[str] = ""
    gstin: Optional[str] = ""
    date: Optional[str] = ""
    valid_until: Optional[str] = ""
    items: List[dict] = []
    discount: Optional[float] = 0
    shipping: Optional[float] = 0
    other_charges: Optional[float] = 0
    terms: Optional[str] = ""
    notes: Optional[str] = ""
    status: Optional[str] = "Draft"


class InvoiceIn(BaseModel):
    customer_id: Optional[str] = ""
    customer_name: str
    company: Optional[str] = ""
    billing_address: Optional[str] = ""
    shipping_address: Optional[str] = ""
    gstin: Optional[str] = ""
    date: Optional[str] = ""
    items: List[dict] = []
    discount: Optional[float] = 0
    shipping: Optional[float] = 0
    other_charges: Optional[float] = 0
    gst_percent: Optional[float] = 18
    payment_status: Optional[str] = "Unpaid"
    payment_method: Optional[str] = ""
    terms: Optional[str] = ""
    notes: Optional[str] = ""


class PaymentIn(BaseModel):
    amount: float
    method: Optional[str] = ""
    date: Optional[str] = ""
    note: Optional[str] = ""


class CatalogueIn(BaseModel):
    name: str
    intro: Optional[str] = ""
    product_ids: List[str] = []
    category_ids: List[str] = []
    template: Optional[str] = "classic"


class BlogIn(BaseModel):
    title: str
    excerpt: Optional[str] = ""
    content: Optional[str] = ""
    cover_image: Optional[str] = ""
    author: Optional[str] = "KEM Enterprises"
    tags: Optional[List[str]] = []
    published: Optional[bool] = True
    seo_title: Optional[str] = ""
    seo_description: Optional[str] = ""


class FaqIn(BaseModel):
    question: str
    answer: str
    order: Optional[int] = 0
    active: Optional[bool] = True


class BulkEditIn(BaseModel):
    ids: List[str] = []
    fields: dict = {}


class BulkDeleteIn(BaseModel):
    ids: List[str] = []
    hard: Optional[bool] = False


class ImportCommitIn(BaseModel):
    rows: List[dict] = []
    filename: Optional[str] = ""


class SettingsIn(BaseModel):
    company: Optional[dict] = None
    contact: Optional[dict] = None
    social: Optional[dict] = None
    whatsapp: Optional[dict] = None
    homepage: Optional[dict] = None
    invoice: Optional[dict] = None
    quotation: Optional[dict] = None
    seo: Optional[dict] = None
    analytics: Optional[dict] = None


def compute_totals(items, discount=0, extra_tax_percent=None, shipping=0, other=0):
    subtotal = 0.0
    tax_total = 0.0
    for it in items:
        qty = float(it.get("quantity", 1) or 0)
        rate = float(it.get("unit_price", 0) or 0)
        d = float(it.get("discount", 0) or 0)
        base = qty * rate * (1 - d / 100)
        subtotal += base
        tax_pct = extra_tax_percent if extra_tax_percent is not None else float(it.get("tax", 0) or 0)
        tax_total += base * (tax_pct / 100)
    after_disc = subtotal - float(discount or 0)
    total = after_disc + tax_total + float(shipping or 0) + float(other or 0)
    return round(subtotal, 2), round(tax_total, 2), round(total, 2)


# ----------------------------- auth -----------------------------
@api.post("/auth/login")
async def login(body: LoginIn, response: Response):
    user = await db.users.find_one({"email": body.email.lower()})
    if not user or not verify_password(body.password, user["password_hash"]):
        raise HTTPException(status_code=401, detail="Invalid email or password")
    token = create_access_token(user["id"], user["email"])
    response.set_cookie("access_token", token, httponly=True, secure=True, samesite="none",
                        max_age=604800, path="/")
    return {"token": token, "user": {"id": user["id"], "name": user["name"],
            "email": user["email"], "role": user["role"],
            "permissions": ROLE_PERMISSIONS.get(user["role"], [])}}


@api.post("/auth/logout")
async def logout(response: Response):
    response.delete_cookie("access_token", path="/")
    return {"ok": True}


@api.get("/auth/me")
async def me(user: dict = Depends(get_current_user)):
    user["permissions"] = ROLE_PERMISSIONS.get(user.get("role"), [])
    return user


# users mgmt
@api.get("/users")
async def list_users(user: dict = Depends(require_section("settings"))):
    users = await db.users.find({}, {"password_hash": 0, "_id": 0}).to_list(200)
    return users


@api.post("/users")
async def create_user(body: UserIn, user: dict = Depends(require_section("settings"))):
    if await db.users.find_one({"email": body.email.lower()}):
        raise HTTPException(status_code=400, detail="Email already exists")
    if body.role not in ROLE_PERMISSIONS:
        raise HTTPException(status_code=400, detail="Invalid role")
    doc = {"id": str(uuid.uuid4()), "name": body.name, "email": body.email.lower(),
           "password_hash": hash_password(body.password), "role": body.role, "created_at": now_iso()}
    await db.users.insert_one(doc)
    return {"id": doc["id"], "name": doc["name"], "email": doc["email"], "role": doc["role"]}


@api.delete("/users/{uid}")
async def delete_user(uid: str, user: dict = Depends(require_section("settings"))):
    target = await db.users.find_one({"id": uid})
    if target and target.get("role") == "super_admin":
        raise HTTPException(status_code=400, detail="Cannot delete super admin")
    await db.users.delete_one({"id": uid})
    return {"ok": True}


# ----------------------------- settings -----------------------------
@api.get("/settings")
async def settings_public():
    return await get_settings()


@api.put("/settings")
async def update_settings(body: SettingsIn, user: dict = Depends(require_section("settings"))):
    update = {k: v for k, v in body.model_dump().items() if v is not None}
    await db.settings.update_one({"id": "global"}, {"$set": update}, upsert=True)
    return await get_settings()


# ----------------------------- categories -----------------------------
async def category_with_count(c):
    c = clean(c)
    c["product_count"] = await db.products.count_documents({"category_id": c["id"], "active": True})
    return c


@api.get("/categories")
async def list_categories(all: bool = False):
    q = {} if all else {"active": True}
    cats = await db.categories.find(q).sort("order", 1).to_list(1000)
    return [await category_with_count(c) for c in cats]


@api.get("/categories/{slug}")
async def get_category(slug: str):
    c = await db.categories.find_one({"slug": slug})
    if not c:
        raise HTTPException(status_code=404, detail="Category not found")
    await log_event("category_view", category_id=c["id"])
    return await category_with_count(c)


@api.post("/categories")
async def create_category(body: CategoryIn, user: dict = Depends(require_section("categories"))):
    slug = await unique_slug(db.categories, slugify(body.name))
    doc = body.model_dump()
    doc.update({"id": str(uuid.uuid4()), "slug": slug, "created_at": now_iso()})
    await db.categories.insert_one(doc)
    return clean(doc)


@api.put("/categories/{cid}")
async def update_category(cid: str, body: CategoryIn, user: dict = Depends(require_section("categories"))):
    existing = await db.categories.find_one({"id": cid})
    if not existing:
        raise HTTPException(status_code=404, detail="Not found")
    update = body.model_dump()
    if body.name != existing.get("name"):
        update["slug"] = await unique_slug(db.categories, slugify(body.name), exclude_id=cid)
    await db.categories.update_one({"id": cid}, {"$set": update})
    return await category_with_count(await db.categories.find_one({"id": cid}))


@api.delete("/categories/{cid}")
async def delete_category(cid: str, user: dict = Depends(require_section("categories"))):
    await db.categories.delete_one({"id": cid})
    return {"ok": True}


@api.post("/categories/reorder")
async def reorder_categories(body: dict, user: dict = Depends(require_section("categories"))):
    for i, cid in enumerate(body.get("order", [])):
        await db.categories.update_one({"id": cid}, {"$set": {"order": i}})
    return {"ok": True}


# ----------------------------- products -----------------------------
@api.get("/products")
async def list_products(category: Optional[str] = None, featured: Optional[bool] = None,
                        is_new: Optional[bool] = None, best_seller: Optional[bool] = None,
                        search: Optional[str] = None, sort: str = "newest",
                        page: int = 1, limit: int = 24, all: bool = False):
    q = {} if all else {"active": True}
    if category:
        cat = await db.categories.find_one({"slug": category})
        if cat:
            q["category_id"] = cat["id"]
    if featured is not None:
        q["featured"] = featured
    if is_new is not None:
        q["is_new"] = is_new
    if best_seller is not None:
        q["best_seller"] = best_seller
    if search:
        rx = {"$regex": search, "$options": "i"}
        q["$or"] = [{"name": rx}, {"sku": rx}, {"category_name": rx}, {"material": rx},
                    {"application": rx}, {"tags": rx}, {"short_description": rx}]
    total = await db.products.count_documents(q)
    sort_map = {"newest": ("created_at", -1), "name_asc": ("name", 1), "name_desc": ("name", -1)}
    sk, sd = sort_map.get(sort, ("created_at", -1))
    cursor = db.products.find(q).sort(sk, sd).skip((page - 1) * limit).limit(limit)
    items = [clean(p) for p in await cursor.to_list(limit)]
    return {"items": items, "total": total, "page": page, "limit": limit}


@api.get("/products/search")
async def search_products(q: str = ""):
    if not q:
        return []
    rx = {"$regex": q, "$options": "i"}
    cursor = db.products.find({"active": True, "$or": [
        {"name": rx}, {"sku": rx}, {"category_name": rx}, {"tags": rx}, {"application": rx}]}).limit(8)
    items = await cursor.to_list(8)
    return [{"name": p["name"], "slug": p["slug"], "category_name": p.get("category_name"),
             "image": (p.get("images") or [""])[0], "sku": p.get("sku")} for p in items]


@api.get("/products/{slug}")
async def get_product(slug: str):
    p = await db.products.find_one({"slug": slug})
    if not p:
        raise HTTPException(status_code=404, detail="Product not found")
    await db.products.update_one({"id": p["id"]}, {"$inc": {"views": 1}})
    await log_event("product_view", product_id=p["id"])
    p = clean(p)
    if p.get("related_ids"):
        related = await db.products.find({"id": {"$in": p["related_ids"]}, "active": True}).to_list(8)
    else:
        related = await db.products.find({"category_id": p["category_id"], "active": True,
                                          "id": {"$ne": p["id"]}}).limit(4).to_list(4)
    p["related"] = [clean(r) for r in related]
    return p


@api.post("/products")
async def create_product(body: ProductIn, user: dict = Depends(require_section("products"))):
    cat = await db.categories.find_one({"id": body.category_id})
    if not cat:
        raise HTTPException(status_code=400, detail="Invalid category")
    slug = await unique_slug(db.products, slugify(body.name))
    doc = body.model_dump()
    doc.update({"id": str(uuid.uuid4()), "slug": slug, "category_name": cat["name"],
                "views": 0, "enquiry_count": 0, "created_at": now_iso()})
    if not doc.get("sku"):
        doc["sku"] = "KEM-" + slug[:10].upper()
    await db.products.insert_one(doc)
    return clean(doc)


@api.put("/products/{pid}")
async def update_product(pid: str, body: ProductIn, user: dict = Depends(require_section("products"))):
    existing = await db.products.find_one({"id": pid})
    if not existing:
        raise HTTPException(status_code=404, detail="Not found")
    cat = await db.categories.find_one({"id": body.category_id})
    update = body.model_dump()
    update["category_name"] = cat["name"] if cat else existing.get("category_name")
    if body.name != existing.get("name"):
        update["slug"] = await unique_slug(db.products, slugify(body.name), exclude_id=pid)
    await db.products.update_one({"id": pid}, {"$set": update})
    return clean(await db.products.find_one({"id": pid}))


@api.post("/products/{pid}/duplicate")
async def duplicate_product(pid: str, user: dict = Depends(require_section("products"))):
    p = await db.products.find_one({"id": pid})
    if not p:
        raise HTTPException(status_code=404, detail="Not found")
    p = clean(dict(p))
    p["id"] = str(uuid.uuid4())
    p["name"] = p["name"] + " (Copy)"
    p["slug"] = await unique_slug(db.products, slugify(p["name"]))
    p["sku"] = (p.get("sku") or "KEM") + "-C"
    p["created_at"] = now_iso()
    p["views"] = 0
    await db.products.insert_one(p)
    return clean(p)


@api.patch("/products/{pid}/flags")
async def toggle_flags(pid: str, body: dict, user: dict = Depends(require_section("products"))):
    allowed = {k: v for k, v in body.items() if k in
               ("featured", "is_new", "best_seller", "popular", "active")}
    await db.products.update_one({"id": pid}, {"$set": allowed})
    return clean(await db.products.find_one({"id": pid}))


@api.delete("/products/{pid}")
async def delete_product(pid: str, user: dict = Depends(require_section("products"))):
    await db.products.delete_one({"id": pid})
    return {"ok": True}


# ----------------------------- media / upload -----------------------------
@api.post("/upload")
async def upload_file(file: UploadFile = File(...), folder: str = Query("general"),
                      user: dict = Depends(get_current_user)):
    ext = (file.filename.rsplit(".", 1)[-1] if "." in file.filename else "bin").lower()
    if ext not in MIME_TYPES:
        raise HTTPException(status_code=400, detail="Unsupported file type")
    path = f"{APP_NAME}/{folder}/{uuid.uuid4()}.{ext}"
    data = await file.read()
    if len(data) > 25 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="File too large (max 25MB)")
    ct = file.content_type or MIME_TYPES.get(ext, "application/octet-stream")
    result = put_object(path, data, ct)
    rec = {"id": str(uuid.uuid4()), "storage_path": result["path"], "original_filename": file.filename,
           "content_type": ct, "size": result.get("size", len(data)), "folder": folder,
           "url": f"/api/files/{result['path']}", "is_deleted": False, "created_at": now_iso()}
    await db.media.insert_one(rec)
    return clean(rec)


@api.get("/files/{path:path}")
async def serve_file(path: str):
    rec = await db.media.find_one({"storage_path": path, "is_deleted": False})
    try:
        data, ct = get_object(path)
    except Exception:
        raise HTTPException(status_code=404, detail="File not found")
    return StarletteResponse(content=data, media_type=(rec or {}).get("content_type", ct))


@api.get("/media")
async def list_media(folder: Optional[str] = None, user: dict = Depends(require_section("media"))):
    q = {"is_deleted": False}
    if folder:
        q["folder"] = folder
    return [clean(m) for m in await db.media.find(q).sort("created_at", -1).to_list(1000)]


@api.delete("/media/{mid}")
async def delete_media(mid: str, user: dict = Depends(require_section("media"))):
    await db.media.update_one({"id": mid}, {"$set": {"is_deleted": True}})
    return {"ok": True}


# ----------------------------- enquiries -----------------------------
@api.post("/enquiries")
async def create_enquiry(body: EnquiryIn):
    doc = body.model_dump()
    doc.update({"id": str(uuid.uuid4()), "status": "New", "assigned_to": "",
                "notes": [], "created_at": now_iso(),
                "month": datetime.now(timezone.utc).strftime("%Y-%m")})
    await db.enquiries.insert_one(doc)
    if body.product_id:
        await db.products.update_one({"id": body.product_id}, {"$inc": {"enquiry_count": 1}})
    await log_event("enquiry", product_id=body.product_id or "")
    return {"ok": True, "id": doc["id"]}


@api.get("/enquiries")
async def list_enquiries(status: Optional[str] = None, user: dict = Depends(require_section("enquiries"))):
    q = {}
    if status:
        q["status"] = status
    return [clean(e) for e in await db.enquiries.find(q).sort("created_at", -1).to_list(1000)]


@api.patch("/enquiries/{eid}")
async def update_enquiry(eid: str, body: EnquiryUpdate, user: dict = Depends(require_section("enquiries"))):
    update = {}
    if body.status:
        update["status"] = body.status
    if body.assigned_to is not None:
        update["assigned_to"] = body.assigned_to
    ops = {}
    if update:
        ops["$set"] = update
    if body.note:
        ops["$push"] = {"notes": {"text": body.note, "by": user["name"], "at": now_iso()}}
    if ops:
        await db.enquiries.update_one({"id": eid}, ops)
    return clean(await db.enquiries.find_one({"id": eid}))


@api.post("/enquiries/{eid}/to-customer")
async def enquiry_to_customer(eid: str, user: dict = Depends(require_section("enquiries"))):
    e = await db.enquiries.find_one({"id": eid})
    if not e:
        raise HTTPException(status_code=404, detail="Not found")
    doc = {"id": str(uuid.uuid4()), "name": e["name"], "company": e.get("company", ""),
           "phone": e.get("phone", ""), "whatsapp": e.get("whatsapp", ""), "email": e.get("email", ""),
           "address": "", "gstin": "", "notes": e.get("requirement", ""), "created_at": now_iso()}
    await db.customers.insert_one(doc)
    return clean(doc)


@api.delete("/enquiries/{eid}")
async def delete_enquiry(eid: str, user: dict = Depends(require_section("enquiries"))):
    await db.enquiries.delete_one({"id": eid})
    return {"ok": True}


# ----------------------------- customers -----------------------------
@api.get("/customers")
async def list_customers(search: Optional[str] = None, user: dict = Depends(require_section("customers"))):
    q = {}
    if search:
        rx = {"$regex": search, "$options": "i"}
        q["$or"] = [{"name": rx}, {"company": rx}, {"phone": rx}, {"email": rx}]
    return [clean(c) for c in await db.customers.find(q).sort("created_at", -1).to_list(1000)]


@api.get("/customers/{cid}")
async def get_customer(cid: str, user: dict = Depends(require_section("customers"))):
    c = await db.customers.find_one({"id": cid})
    if not c:
        raise HTTPException(status_code=404, detail="Not found")
    c = clean(c)
    c["quotations"] = [clean(x) for x in await db.quotations.find({"customer_id": cid}).to_list(100)]
    c["invoices"] = [clean(x) for x in await db.invoices.find({"customer_id": cid}).to_list(100)]
    return c


@api.post("/customers")
async def create_customer(body: CustomerIn, user: dict = Depends(require_section("customers"))):
    doc = body.model_dump()
    doc.update({"id": str(uuid.uuid4()), "created_at": now_iso()})
    await db.customers.insert_one(doc)
    return clean(doc)


@api.put("/customers/{cid}")
async def update_customer(cid: str, body: CustomerIn, user: dict = Depends(require_section("customers"))):
    await db.customers.update_one({"id": cid}, {"$set": body.model_dump()})
    return clean(await db.customers.find_one({"id": cid}))


@api.delete("/customers/{cid}")
async def delete_customer(cid: str, user: dict = Depends(require_section("customers"))):
    await db.customers.delete_one({"id": cid})
    return {"ok": True}


# ----------------------------- quotations -----------------------------
@api.get("/quotations")
async def list_quotations(user: dict = Depends(require_section("quotations"))):
    return [clean(q) for q in await db.quotations.find({}).sort("created_at", -1).to_list(1000)]


@api.get("/quotations/{qid}")
async def get_quotation(qid: str, user: dict = Depends(require_section("quotations"))):
    q = await db.quotations.find_one({"id": qid})
    if not q:
        raise HTTPException(status_code=404, detail="Not found")
    return clean(q)


@api.post("/quotations")
async def create_quotation(body: QuotationIn, user: dict = Depends(require_section("quotations"))):
    subtotal, tax, total = compute_totals(body.items, body.discount, None, body.shipping, body.other_charges)
    doc = body.model_dump()
    doc.update({"id": str(uuid.uuid4()), "number": await next_number("quotation"),
                "date": body.date or now_iso()[:10], "subtotal": subtotal, "tax": tax, "total": total,
                "created_at": now_iso(), "month": datetime.now(timezone.utc).strftime("%Y-%m")})
    if not body.valid_until:
        s = await get_settings()
        days = s.get("quotation", {}).get("validity", 15)
        doc["valid_until"] = (datetime.now(timezone.utc) + timedelta(days=days)).isoformat()[:10]
    if not body.terms:
        doc["terms"] = (await get_settings()).get("quotation", {}).get("terms", "")
    await db.quotations.insert_one(doc)
    return clean(doc)


@api.put("/quotations/{qid}")
async def update_quotation(qid: str, body: QuotationIn, user: dict = Depends(require_section("quotations"))):
    subtotal, tax, total = compute_totals(body.items, body.discount, None, body.shipping, body.other_charges)
    update = body.model_dump()
    update.update({"subtotal": subtotal, "tax": tax, "total": total})
    await db.quotations.update_one({"id": qid}, {"$set": update})
    return clean(await db.quotations.find_one({"id": qid}))


@api.delete("/quotations/{qid}")
async def delete_quotation(qid: str, user: dict = Depends(require_section("quotations"))):
    await db.quotations.delete_one({"id": qid})
    return {"ok": True}


@api.get("/quotations/{qid}/pdf")
async def quotation_pdf(qid: str, user: dict = Depends(require_section("quotations"))):
    q = clean(await db.quotations.find_one({"id": qid}))
    if not q:
        raise HTTPException(status_code=404, detail="Not found")
    s = await get_settings()
    party = {"name": q.get("customer_name"), "company": q.get("company"),
             "address": q.get("address"), "gstin": q.get("gstin")}
    pdf = pdfgen.build_document_pdf("quotation", q, s.get("company", {}), party)
    return StarletteResponse(content=pdf, media_type="application/pdf",
                             headers={"Content-Disposition": f'inline; filename="{q["number"]}.pdf"'})


@api.post("/quotations/{qid}/convert-invoice")
async def convert_to_invoice(qid: str, user: dict = Depends(require_section("invoices"))):
    q = await db.quotations.find_one({"id": qid})
    if not q:
        raise HTTPException(status_code=404, detail="Not found")
    s = await get_settings()
    gst = s.get("invoice", {}).get("gst", 18)
    subtotal, tax, total = compute_totals(q["items"], q.get("discount", 0), gst,
                                          q.get("shipping", 0), q.get("other_charges", 0))
    doc = {"id": str(uuid.uuid4()), "number": await next_number("invoice"),
           "customer_id": q.get("customer_id", ""), "customer_name": q["customer_name"],
           "company": q.get("company", ""), "billing_address": q.get("address", ""),
           "shipping_address": q.get("address", ""), "gstin": q.get("gstin", ""),
           "date": now_iso()[:10], "items": q["items"], "discount": q.get("discount", 0),
           "shipping": q.get("shipping", 0), "other_charges": q.get("other_charges", 0),
           "gst_percent": gst, "subtotal": subtotal, "gst": tax, "total": total,
           "payment_status": "Unpaid", "payment_method": "", "payments": [],
           "terms": s.get("invoice", {}).get("terms", ""), "notes": "",
           "created_at": now_iso(), "month": datetime.now(timezone.utc).strftime("%Y-%m")}
    await db.invoices.insert_one(doc)
    await db.quotations.update_one({"id": qid}, {"$set": {"status": "Won"}})
    return clean(doc)


# ----------------------------- invoices -----------------------------
@api.get("/invoices")
async def list_invoices(user: dict = Depends(require_section("invoices"))):
    return [clean(i) for i in await db.invoices.find({}).sort("created_at", -1).to_list(1000)]


@api.get("/invoices/{iid}")
async def get_invoice(iid: str, user: dict = Depends(require_section("invoices"))):
    i = await db.invoices.find_one({"id": iid})
    if not i:
        raise HTTPException(status_code=404, detail="Not found")
    return clean(i)


@api.post("/invoices")
async def create_invoice(body: InvoiceIn, user: dict = Depends(require_section("invoices"))):
    subtotal, tax, total = compute_totals(body.items, body.discount, body.gst_percent,
                                          body.shipping, body.other_charges)
    doc = body.model_dump()
    doc.update({"id": str(uuid.uuid4()), "number": await next_number("invoice"),
                "date": body.date or now_iso()[:10], "subtotal": subtotal, "gst": tax, "total": total,
                "payments": [], "created_at": now_iso(),
                "month": datetime.now(timezone.utc).strftime("%Y-%m")})
    if not body.terms:
        doc["terms"] = (await get_settings()).get("invoice", {}).get("terms", "")
    await db.invoices.insert_one(doc)
    return clean(doc)


@api.put("/invoices/{iid}")
async def update_invoice(iid: str, body: InvoiceIn, user: dict = Depends(require_section("invoices"))):
    subtotal, tax, total = compute_totals(body.items, body.discount, body.gst_percent,
                                          body.shipping, body.other_charges)
    update = body.model_dump()
    update.update({"subtotal": subtotal, "gst": tax, "total": total})
    await db.invoices.update_one({"id": iid}, {"$set": update})
    return clean(await db.invoices.find_one({"id": iid}))


@api.post("/invoices/{iid}/payments")
async def record_payment(iid: str, body: PaymentIn, user: dict = Depends(require_section("payments"))):
    inv = await db.invoices.find_one({"id": iid})
    if not inv:
        raise HTTPException(status_code=404, detail="Not found")
    payment = {"id": str(uuid.uuid4()), "amount": body.amount, "method": body.method,
               "date": body.date or now_iso()[:10], "note": body.note}
    payments = inv.get("payments", []) + [payment]
    paid = sum(p["amount"] for p in payments)
    status = "Paid" if paid >= inv["total"] else ("Partially Paid" if paid > 0 else "Unpaid")
    await db.invoices.update_one({"id": iid}, {"$set": {"payments": payments, "payment_status": status,
                                                        "payment_method": body.method}})
    return clean(await db.invoices.find_one({"id": iid}))


@api.delete("/invoices/{iid}")
async def delete_invoice(iid: str, user: dict = Depends(require_section("invoices"))):
    await db.invoices.delete_one({"id": iid})
    return {"ok": True}


@api.get("/invoices/{iid}/pdf")
async def invoice_pdf(iid: str, user: dict = Depends(require_section("invoices"))):
    inv = clean(await db.invoices.find_one({"id": iid}))
    if not inv:
        raise HTTPException(status_code=404, detail="Not found")
    s = await get_settings()
    party = {"name": inv.get("customer_name"), "company": inv.get("company"),
             "address": inv.get("billing_address"), "gstin": inv.get("gstin")}
    pdf = pdfgen.build_document_pdf("invoice", inv, s.get("company", {}), party)
    return StarletteResponse(content=pdf, media_type="application/pdf",
                             headers={"Content-Disposition": f'inline; filename="{inv["number"]}.pdf"'})


# ----------------------------- catalogues -----------------------------
async def catalogue_products(cat):
    ordered = list(cat.get("product_ids", []))
    if cat.get("category_ids"):
        cprods = await db.products.find({"category_id": {"$in": cat["category_ids"]}, "active": True}).to_list(1000)
        for p in cprods:
            if p["id"] not in ordered:
                ordered.append(p["id"])
    if not ordered:
        prods = await db.products.find({"active": True}).to_list(1000)
        return [clean(p) for p in prods]
    found = {p["id"]: clean(p) for p in await db.products.find({"id": {"$in": ordered}}).to_list(1000)}
    return [found[i] for i in ordered if i in found]


@api.get("/catalogues")
async def list_catalogues(user: dict = Depends(require_section("catalogues"))):
    return [clean(c) for c in await db.catalogues.find({}).sort("created_at", -1).to_list(200)]


@api.post("/catalogues")
async def create_catalogue(body: CatalogueIn, user: dict = Depends(require_section("catalogues"))):
    doc = body.model_dump()
    doc.update({"id": str(uuid.uuid4()), "share_token": uuid.uuid4().hex[:12],
                "downloads": 0, "opens": 0, "shares": 0, "created_at": now_iso()})
    await db.catalogues.insert_one(doc)
    return clean(doc)


@api.put("/catalogues/{cid}")
async def update_catalogue(cid: str, body: CatalogueIn, user: dict = Depends(require_section("catalogues"))):
    await db.catalogues.update_one({"id": cid}, {"$set": body.model_dump()})
    return clean(await db.catalogues.find_one({"id": cid}))


@api.delete("/catalogues/{cid}")
async def delete_catalogue(cid: str, user: dict = Depends(require_section("catalogues"))):
    await db.catalogues.delete_one({"id": cid})
    return {"ok": True}


async def _build_catalogue_pdf(cat):
    s = await get_settings()
    prods = await catalogue_products(cat)
    backend = os.environ.get("REACT_APP_BACKEND_URL", "")
    wa = s.get("whatsapp", {}).get("number", "")
    return pdfgen.build_catalogue_pdf(cat, prods, s.get("company", {}),
                                      backend or "https://kem.example.com",
                                      f"https://wa.me/{wa}")


@api.get("/catalogue-full/pdf")
async def catalogue_full_pdf():
    s = await get_settings()
    prods = [clean(p) for p in await db.products.find({"active": True}).sort("category_name", 1).to_list(2000)]
    backend = os.environ.get("REACT_APP_BACKEND_URL", "")
    wa = s.get("whatsapp", {}).get("number", "")
    cat = {"name": "Complete Product Catalogue", "intro": s.get("company", {}).get("about", ""), "template": "classic"}
    pdf = pdfgen.build_catalogue_pdf(cat, prods, s.get("company", {}),
                                     backend or "https://kem.example.com", f"https://wa.me/{wa}")
    return StarletteResponse(content=pdf, media_type="application/pdf",
                             headers={"Content-Disposition": 'inline; filename="kem-complete-catalogue.pdf"'})


@api.get("/catalogues/{cid}/pdf")
async def catalogue_pdf(cid: str, user: dict = Depends(require_section("catalogues"))):
    cat = clean(await db.catalogues.find_one({"id": cid}))
    if not cat:
        raise HTTPException(status_code=404, detail="Not found")
    await db.catalogues.update_one({"id": cid}, {"$inc": {"downloads": 1}})
    pdf = await _build_catalogue_pdf(cat)
    return StarletteResponse(content=pdf, media_type="application/pdf",
                             headers={"Content-Disposition": f'inline; filename="{slugify(cat["name"])}.pdf"'})


@api.post("/catalogues/{cid}/share")
async def catalogue_share(cid: str, body: dict, user: dict = Depends(require_section("catalogues"))):
    cat = await db.catalogues.find_one({"id": cid})
    if not cat:
        raise HTTPException(status_code=404, detail="Not found")
    await db.catalogues.update_one({"id": cid}, {"$inc": {"shares": 1}})
    await db.catalogue_shares.insert_one({"id": str(uuid.uuid4()), "catalogue_id": cid,
                                          "channel": body.get("channel", ""), "to": body.get("to", ""),
                                          "created_at": now_iso()})
    backend = os.environ.get("REACT_APP_BACKEND_URL", "")
    return {"link": f"{backend}/catalogue/view/{cat['share_token']}"}


@api.get("/public/catalogue/{token}")
async def public_catalogue(token: str):
    cat = await db.catalogues.find_one({"share_token": token})
    if not cat:
        raise HTTPException(status_code=404, detail="Catalogue not found")
    await db.catalogues.update_one({"id": cat["id"]}, {"$inc": {"opens": 1}})
    cat = clean(cat)
    cat["products"] = await catalogue_products(cat)
    return cat


@api.get("/public/catalogue/{token}/pdf")
async def public_catalogue_pdf(token: str):
    cat = clean(await db.catalogues.find_one({"share_token": token}))
    if not cat:
        raise HTTPException(status_code=404, detail="Not found")
    await db.catalogues.update_one({"id": cat["id"]}, {"$inc": {"downloads": 1}})
    pdf = await _build_catalogue_pdf(cat)
    return StarletteResponse(content=pdf, media_type="application/pdf",
                             headers={"Content-Disposition": f'inline; filename="{slugify(cat["name"])}.pdf"'})


# ----------------------------- product compare / import / export -----------------------------
@api.post("/compare")
async def compare_products(body: dict):
    ids = (body.get("ids") or [])[:4]
    prods = await db.products.find({"id": {"$in": ids}, "active": True}).to_list(4)
    order = {pid: i for i, pid in enumerate(ids)}
    prods.sort(key=lambda p: order.get(p["id"], 99))
    return [clean(p) for p in prods]


@api.get("/products-export")
async def export_products(user: dict = Depends(require_section("products"))):
    prods = await db.products.find({}).to_list(5000)
    fields = ["name", "sku", "category_name", "short_description", "description", "material",
              "finish", "size", "load_capacity", "application", "moq", "packaging", "tags",
              "images", "featured", "is_new", "best_seller", "active"]
    out = io.StringIO()
    w = csv.DictWriter(out, fieldnames=fields)
    w.writeheader()
    for p in prods:
        w.writerow({f: (",".join(p.get(f) or []) if isinstance(p.get(f), list) else p.get(f, "")) for f in fields})
    return StarletteResponse(content=out.getvalue(), media_type="text/csv",
                             headers={"Content-Disposition": 'attachment; filename="kem-products.csv"'})


@api.post("/products-import")
async def import_products(file: UploadFile = File(...), user: dict = Depends(require_section("products"))):
    data = await file.read()
    ext = (file.filename.rsplit(".", 1)[-1] if "." in file.filename else "csv").lower()
    rows = []
    if ext in ("xlsx", "xls"):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True)
        ws = wb.active
        it = ws.iter_rows(values_only=True)
        headers = [str(h).strip() if h is not None else "" for h in next(it)]
        for r in it:
            rows.append({headers[i]: (r[i] if i < len(r) and r[i] is not None else "") for i in range(len(headers))})
    else:
        rows = list(csv.DictReader(io.StringIO(data.decode("utf-8-sig"))))

    cats = await db.categories.find({}).to_list(1000)
    cat_by_name = {c["name"].strip().lower(): c for c in cats}
    created, updated = 0, 0

    def b(v):
        return str(v).strip().lower() in ("1", "true", "yes", "y")

    for row in rows:
        name = str(row.get("name") or "").strip()
        if not name:
            continue
        cat_name = str(row.get("category_name") or "").strip() or "New Items"
        cat = cat_by_name.get(cat_name.lower())
        if not cat:
            slug = await unique_slug(db.categories, slugify(cat_name))
            cat = {"id": str(uuid.uuid4()), "name": cat_name, "slug": slug, "description": "",
                   "image": "", "parent_id": None, "order": 999, "active": True,
                   "seo_title": "", "seo_description": "", "seo_keywords": "", "created_at": now_iso()}
            await db.categories.insert_one(dict(cat))
            cat_by_name[cat_name.lower()] = cat
        tags = [t.strip() for t in str(row.get("tags") or "").split(",") if t.strip()]
        images = [t.strip() for t in str(row.get("images") or "").split(",") if t.strip()]
        active_raw = str(row.get("active") or "").strip()
        fields = {
            "category_id": cat["id"], "category_name": cat["name"],
            "short_description": str(row.get("short_description") or ""),
            "description": str(row.get("description") or row.get("short_description") or ""),
            "material": str(row.get("material") or ""), "finish": str(row.get("finish") or ""),
            "size": str(row.get("size") or ""), "load_capacity": str(row.get("load_capacity") or ""),
            "application": str(row.get("application") or ""), "moq": str(row.get("moq") or ""),
            "packaging": str(row.get("packaging") or ""), "tags": tags, "images": images,
            "featured": b(row.get("featured")), "is_new": b(row.get("is_new")),
            "best_seller": b(row.get("best_seller")), "active": (b(active_raw) if active_raw else True),
        }
        sku = str(row.get("sku") or "").strip()
        existing = (await db.products.find_one({"sku": sku})) if sku else None
        if not existing:
            existing = await db.products.find_one({"name": name})
        if existing:
            await db.products.update_one({"id": existing["id"]}, {"$set": fields})
            updated += 1
        else:
            slug = await unique_slug(db.products, slugify(name))
            doc = {"id": str(uuid.uuid4()), "name": name, "slug": slug,
                   "sku": sku or ("KEM-" + slug[:10].upper()), "video": "", "specifications": [],
                   "dimensions": fields["size"], "wire_diameter": "", "thread_size": "",
                   "variants": [], "brand": "KEM Enterprises", "related_ids": [], "downloads": [],
                   "price": 0, "show_price": False, "popular": False, "views": 0, "enquiry_count": 0,
                   "subcategory": "", "seo_title": name, "seo_description": "", "seo_keywords": ",".join(tags),
                   "created_at": now_iso(), **fields}
            await db.products.insert_one(doc)
            created += 1
    return {"created": created, "updated": updated, "total": len([r for r in rows if str(r.get("name") or "").strip()])}


# ----------------------------- bulk import / bulk ops / templates -----------------------------
FIELD_SYNONYMS = {
    "name": ["product name", "product", "product title", "item name", "name", "title"],
    "sku": ["sku", "product code", "code", "sku / product code", "sku/product code", "item code"],
    "category_name": ["category", "category name"],
    "subcategory": ["subcategory", "sub category", "sub-category"],
    "short_description": ["short description", "short desc", "summary"],
    "description": ["description", "full description", "detailed description", "details"],
    "material": ["material"], "finish": ["finish"], "size": ["size"],
    "dimensions": ["dimensions", "dimension"], "wire_diameter": ["wire diameter", "wire dia", "diameter"],
    "load_capacity": ["load capacity", "load", "capacity"],
    "application": ["application", "applications", "use"],
    "moq": ["moq", "minimum order", "min order"], "brand": ["brand"],
    "tags": ["tags", "keywords", "tag"],
    "featured": ["featured", "featured product"], "is_new": ["new", "new product", "is new"],
    "best_seller": ["best seller", "bestseller"], "active": ["status", "active", "product status"],
    "seo_title": ["seo title"], "seo_description": ["seo description", "meta description"],
    "seo_keywords": ["seo keywords", "meta keywords"],
    "images": ["product image", "image", "image url", "images", "photo"],
    "datasheet": ["datasheet pdf", "datasheet", "pdf", "pdf url", "datasheet url"],
}
PRODUCT_TEMPLATE_COLS = [
    ("Product Name", True), ("SKU / Product Code", False), ("Category", True), ("Subcategory", False),
    ("Short Description", False), ("Full Description", False), ("Material", False), ("Finish", False),
    ("Size", False), ("Dimensions", False), ("Wire Diameter", False), ("Load Capacity", False),
    ("Application", False), ("MOQ", False), ("Brand", False), ("Tags", False), ("Featured Product", False),
    ("New Product", False), ("Product Status", False), ("SEO Title", False), ("SEO Description", False),
    ("SEO Keywords", False), ("Product Image", False), ("Image URL", False), ("Datasheet PDF", False),
]
CATEGORY_TEMPLATE_COLS = ["Category Name", "Parent Category", "Description", "Image", "SEO Title", "SEO Description", "SEO Keywords", "Slug", "Status"]


def _truthy(v):
    return str(v).strip().lower() in ("1", "true", "yes", "y", "active", "enabled")


def _parse_rows(data, ext):
    if ext in ("xlsx", "xls"):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.active
        it = ws.iter_rows(values_only=True)
        headers = [str(h).strip() if h is not None else "" for h in next(it)]
        rows = []
        for r in it:
            if all(c is None or str(c).strip() == "" for c in r):
                continue
            rows.append({headers[i]: (r[i] if i < len(r) and r[i] is not None else "") for i in range(len(headers))})
        return headers, rows
    reader = csv.DictReader(io.StringIO(data.decode("utf-8-sig")))
    return (reader.fieldnames or []), [dict(r) for r in reader]


def _map_headers(headers):
    mapping = {}
    for h in headers:
        hl = (h or "").strip().lower()
        for field, syns in FIELD_SYNONYMS.items():
            if hl in syns and field not in mapping.values():
                mapping[h] = field
                break
    return mapping


def _normalize(row, mapping):
    out = {}
    imgs = []
    for h, field in mapping.items():
        val = row.get(h, "")
        val = "" if val is None else str(val).strip()
        if field == "images":
            if val:
                imgs += [x.strip() for x in val.split(",") if x.strip()]
        else:
            out[field] = val
    out["images"] = imgs
    return out


@api.post("/products/import-preview")
async def import_preview(file: UploadFile = File(...), user: dict = Depends(require_section("products"))):
    data = await file.read()
    ext = (file.filename.rsplit(".", 1)[-1] if "." in file.filename else "csv").lower()
    if ext not in ("csv", "xlsx", "xls"):
        raise HTTPException(status_code=400, detail="Only .csv and .xlsx are supported for structured import")
    headers, rows = _parse_rows(data, ext)
    mapping = _map_headers(headers)
    cats = {c["name"].strip().lower() for c in await db.categories.find({}, {"name": 1}).to_list(1000)}
    out, ready, dup, errors = [], 0, 0, 0
    for i, r in enumerate(rows):
        d = _normalize(r, mapping)
        name = d.get("name", "")
        cat = d.get("category_name", "")
        issues, status = [], "ready"
        if not name:
            issues.append("Missing product name"); status = "error"
        if not cat:
            issues.append("Missing category"); status = "error"
        duplicate_id = ""
        if name:
            existing = await db.products.find_one({"$or": [{"sku": d.get("sku", "x-none")}, {"name": name}, {"slug": slugify(name)}]}) if (d.get("sku") or name) else None
            if existing:
                duplicate_id = existing["id"]
                if status != "error":
                    status = "duplicate"; issues.append("Possible duplicate — will update existing")
        if cat and cat.lower() not in cats and status != "error":
            issues.append("New category will be created")
        action = "skip" if status == "error" else ("update" if status == "duplicate" else "create")
        if status == "ready":
            ready += 1
        elif status == "duplicate":
            dup += 1
        else:
            errors += 1
        out.append({"row": i + 2, "data": d, "status": status, "issues": issues, "action": action, "duplicate_id": duplicate_id})
    return {"mapping": mapping, "headers": headers, "rows": out,
            "summary": {"total": len(rows), "ready": ready, "duplicates": dup, "errors": errors}}


async def _upsert_product(d, action):
    name = (d.get("name") or "").strip()
    if not name or action == "skip":
        return None
    cat_name = (d.get("category_name") or "").strip() or "New Items"
    cat = await db.categories.find_one({"name": {"$regex": f"^{re.escape(cat_name)}$", "$options": "i"}})
    if not cat:
        slug = await unique_slug(db.categories, slugify(cat_name))
        cat = {"id": str(uuid.uuid4()), "name": cat_name, "slug": slug, "description": "", "image": "",
               "parent_id": None, "order": 999, "active": True, "seo_title": "", "seo_description": "",
               "seo_keywords": "", "created_at": now_iso()}
        await db.categories.insert_one(dict(cat))
    tags = [t.strip() for t in str(d.get("tags") or "").split(",") if t.strip()]
    images = d.get("images") or []
    fields = {"category_id": cat["id"], "category_name": cat["name"], "subcategory": d.get("subcategory", ""),
              "short_description": d.get("short_description", ""), "description": d.get("description") or d.get("short_description", ""),
              "material": d.get("material", ""), "finish": d.get("finish", ""), "size": d.get("size", ""),
              "dimensions": d.get("dimensions") or d.get("size", ""), "wire_diameter": d.get("wire_diameter", ""),
              "load_capacity": d.get("load_capacity", ""), "application": d.get("application", ""),
              "moq": d.get("moq", ""), "brand": d.get("brand") or "KEM Enterprises", "tags": tags, "images": images,
              "featured": _truthy(d.get("featured")), "is_new": _truthy(d.get("is_new")),
              "best_seller": _truthy(d.get("best_seller")),
              "active": (_truthy(d.get("active")) if str(d.get("active") or "").strip() else True),
              "seo_title": d.get("seo_title") or name, "seo_description": d.get("seo_description", ""),
              "seo_keywords": d.get("seo_keywords") or ",".join(tags)}
    if d.get("datasheet"):
        fields["downloads"] = [{"name": "Datasheet", "url": d["datasheet"]}]
    sku = (d.get("sku") or "").strip()
    existing = (await db.products.find_one({"sku": sku})) if sku else None
    if not existing:
        existing = await db.products.find_one({"name": name})
    if existing and action != "create":
        await db.products.update_one({"id": existing["id"]}, {"$set": fields})
        return "updated"
    slug = await unique_slug(db.products, slugify(name))
    doc = {"id": str(uuid.uuid4()), "name": name, "slug": slug, "sku": sku or ("KEM-" + slug[:10].upper()),
           "video": "", "specifications": [], "thread_size": "", "variants": [], "related_ids": [],
           "downloads": fields.pop("downloads", []), "price": 0, "show_price": False, "popular": False,
           "views": 0, "enquiry_count": 0, "archived": False, "created_at": now_iso(), **fields}
    await db.products.insert_one(doc)
    return "created"


@api.post("/products/import-commit")
async def import_commit(body: ImportCommitIn, user: dict = Depends(require_section("products"))):
    created = updated = skipped = failed = 0
    errors = []
    for row in body.rows:
        d = row.get("data", row)
        action = row.get("action", "create")
        try:
            res = await _upsert_product(d, action)
            if res == "created":
                created += 1
            elif res == "updated":
                updated += 1
            else:
                skipped += 1
        except Exception as e:
            failed += 1
            errors.append({"row": row.get("row"), "error": str(e)})
    log = {"id": str(uuid.uuid4()), "filename": body.filename, "by": user.get("name"), "type": "products",
           "total": len(body.rows), "created": created, "updated": updated, "skipped": skipped,
           "failed": failed, "errors": errors, "created_at": now_iso()}
    await db.import_logs.insert_one(dict(log))
    return clean(log)


@api.get("/import-history")
async def import_history(user: dict = Depends(require_section("products"))):
    return [clean(x) for x in await db.import_logs.find({}).sort("created_at", -1).to_list(100)]


@api.post("/products/bulk-delete")
async def bulk_delete(body: BulkDeleteIn, user: dict = Depends(require_section("products"))):
    if body.hard:
        r = await db.products.delete_many({"id": {"$in": body.ids}})
        return {"deleted": r.deleted_count}
    r = await db.products.update_many({"id": {"$in": body.ids}}, {"$set": {"active": False, "archived": True}})
    return {"archived": r.modified_count}


@api.post("/products/bulk-edit")
async def bulk_edit(body: BulkEditIn, user: dict = Depends(require_section("products"))):
    allowed = {"category_id", "subcategory", "brand", "active", "featured", "is_new", "best_seller", "popular", "tags"}
    fields = {k: v for k, v in body.fields.items() if k in allowed and v not in (None, "")}
    if "category_id" in fields:
        cat = await db.categories.find_one({"id": fields["category_id"]})
        if cat:
            fields["category_name"] = cat["name"]
    if not fields:
        raise HTTPException(status_code=400, detail="No valid fields to update")
    r = await db.products.update_many({"id": {"$in": body.ids}}, {"$set": fields})
    return {"updated": r.modified_count, "fields": list(fields.keys())}


@api.post("/categories/import-commit")
async def categories_import(body: ImportCommitIn, user: dict = Depends(require_section("categories"))):
    created = updated = 0
    for row in body.rows:
        d = row.get("data", row)
        name = (d.get("Category Name") or d.get("name") or d.get("category_name") or "").strip()
        if not name:
            continue
        existing = await db.categories.find_one({"name": {"$regex": f"^{re.escape(name)}$", "$options": "i"}})
        fields = {"description": d.get("Description") or d.get("description", ""), "image": d.get("Image") or d.get("image", ""),
                  "seo_title": d.get("SEO Title") or "", "seo_description": d.get("SEO Description") or "",
                  "seo_keywords": d.get("SEO Keywords") or "",
                  "active": (_truthy(d.get("Status")) if str(d.get("Status") or "").strip() else True)}
        if existing:
            await db.categories.update_one({"id": existing["id"]}, {"$set": fields})
            updated += 1
        else:
            slug = await unique_slug(db.categories, slugify(name))
            await db.categories.insert_one({"id": str(uuid.uuid4()), "name": name, "slug": slug,
                                            "parent_id": None, "order": 999, "created_at": now_iso(), **fields})
            created += 1
    return {"created": created, "updated": updated, "total": len(body.rows)}


def _xlsx_template(columns, sheet_name):
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    hdr_fill = PatternFill("solid", fgColor="0F172A")
    req_fill = PatternFill("solid", fgColor="FF4D00")
    cols = columns if isinstance(columns[0], tuple) else [(c, False) for c in columns]
    for i, (col, required) in enumerate(cols, 1):
        cell = ws.cell(row=1, column=i, value=col + (" *" if required else ""))
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = req_fill if required else hdr_fill
        ws.column_dimensions[cell.column_letter].width = max(16, len(col) + 4)
    ws.cell(row=2, column=1, value="(example) Sample Product").font = Font(italic=True, color="94A3B8")
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


@api.get("/templates/products.xlsx")
async def products_template():
    data = _xlsx_template(PRODUCT_TEMPLATE_COLS, "Products")
    return StarletteResponse(content=data, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": 'attachment; filename="KEM_Product_Template.xlsx"'})


@api.get("/templates/categories.xlsx")
async def categories_template():
    data = _xlsx_template(CATEGORY_TEMPLATE_COLS, "Categories")
    return StarletteResponse(content=data, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": 'attachment; filename="KEM_Category_Template.xlsx"'})


@api.get("/sitemap.xml")
async def sitemap():
    # PUBLIC_SITE_URL is the canonical public website origin (not the API/backend origin).
    # Keep the example placeholder until the production domain is purchased.
    from xml.sax.saxutils import escape
    base = os.environ.get("PUBLIC_SITE_URL", "https://www.example.com").rstrip("/")
    static_urls = ["/", "/products", "/categories", "/new-products", "/featured-products", "/catalogue",
                   "/about", "/services", "/industries", "/gallery", "/contact", "/blog", "/faq", "/privacy", "/terms"]
    cats = await db.categories.find({"active": True}, {"id": 1, "slug": 1}).to_list(1000)
    cat_by_id = {c["id"]: c["slug"] for c in cats}
    prods = await db.products.find({"active": True, "archived": {"$ne": True}}, {"slug": 1, "category_id": 1}).to_list(5000)
    blogs = await db.blog_posts.find({"published": True}, {"slug": 1}).to_list(500)
    entries = [f"{base}{u}" for u in static_urls]
    entries += [f"{base}/products/{c['slug']}" for c in cats]
    entries += [f"{base}/products/{cat_by_id[p['category_id']]}/{p['slug']}" for p in prods if p.get("category_id") in cat_by_id]
    entries += [f"{base}/blog/{b['slug']}" for b in blogs]
    body = '<?xml version="1.0" encoding="UTF-8"?>\n<urlset xmlns="http://www.sitemaps.org/schemas/sitemap/0.9">\n'
    body += "\n".join(f"  <url><loc>{escape(u)}</loc></url>" for u in entries)
    body += "\n</urlset>"
    return StarletteResponse(content=body, media_type="application/xml")


@api.get("/robots.txt")
async def robots():
    base = os.environ.get("PUBLIC_SITE_URL", "https://www.example.com").rstrip("/")
    body = ("User-agent: *\nAllow: /\nDisallow: /admin\nDisallow: /search\nDisallow: /compare\n"
            "Disallow: /enquiry\nDisallow: /request-quote\nDisallow: /catalogue/view/\nDisallow: /api/\n"
            f"Sitemap: {base}/sitemap.xml\n")
    return StarletteResponse(content=body, media_type="text/plain")


# ----------------------------- blog / knowledge centre -----------------------------
@api.get("/blog")
async def list_blog():
    posts = await db.blog_posts.find({"published": True}).sort("created_at", -1).to_list(200)
    return [clean(p) for p in posts]


@api.get("/blog-admin")
async def list_blog_admin(user: dict = Depends(require_section("content"))):
    return [clean(p) for p in await db.blog_posts.find({}).sort("created_at", -1).to_list(500)]


@api.get("/blog/{slug}")
async def get_blog(slug: str):
    p = await db.blog_posts.find_one({"slug": slug})
    if not p:
        raise HTTPException(status_code=404, detail="Post not found")
    return clean(p)


@api.post("/blog")
async def create_blog(body: BlogIn, user: dict = Depends(require_section("content"))):
    slug = await unique_slug(db.blog_posts, slugify(body.title))
    doc = body.model_dump()
    doc.update({"id": str(uuid.uuid4()), "slug": slug, "created_at": now_iso()})
    await db.blog_posts.insert_one(doc)
    return clean(doc)


@api.put("/blog/{pid}")
async def update_blog(pid: str, body: BlogIn, user: dict = Depends(require_section("content"))):
    existing = await db.blog_posts.find_one({"id": pid})
    if not existing:
        raise HTTPException(status_code=404, detail="Not found")
    update = body.model_dump()
    if body.title != existing.get("title"):
        update["slug"] = await unique_slug(db.blog_posts, slugify(body.title), exclude_id=pid)
    await db.blog_posts.update_one({"id": pid}, {"$set": update})
    return clean(await db.blog_posts.find_one({"id": pid}))


@api.delete("/blog/{pid}")
async def delete_blog(pid: str, user: dict = Depends(require_section("content"))):
    await db.blog_posts.delete_one({"id": pid})
    return {"ok": True}


# ----------------------------- FAQs -----------------------------
@api.get("/faqs")
async def list_faqs():
    return [clean(f) for f in await db.faqs.find({"active": True}).sort("order", 1).to_list(200)]


@api.get("/faqs-admin")
async def list_faqs_admin(user: dict = Depends(require_section("content"))):
    return [clean(f) for f in await db.faqs.find({}).sort("order", 1).to_list(500)]


@api.post("/faqs")
async def create_faq(body: FaqIn, user: dict = Depends(require_section("content"))):
    doc = body.model_dump()
    doc.update({"id": str(uuid.uuid4()), "created_at": now_iso()})
    await db.faqs.insert_one(doc)
    return clean(doc)


@api.put("/faqs/{fid}")
async def update_faq(fid: str, body: FaqIn, user: dict = Depends(require_section("content"))):
    await db.faqs.update_one({"id": fid}, {"$set": body.model_dump()})
    return clean(await db.faqs.find_one({"id": fid}))


@api.delete("/faqs/{fid}")
async def delete_faq(fid: str, user: dict = Depends(require_section("content"))):
    await db.faqs.delete_one({"id": fid})
    return {"ok": True}


# ----------------------------- dashboard / reports / analytics -----------------------------
@api.post("/track")
async def track(body: dict):
    await log_event(body.get("event", "page_view"), path=body.get("path", ""), ref=body.get("ref", ""))
    return {"ok": True}


@api.get("/dashboard/stats")
async def dashboard_stats(user: dict = Depends(get_current_user)):
    total_products = await db.products.count_documents({})
    active_products = await db.products.count_documents({"active": True})
    categories = await db.categories.count_documents({})
    enquiries = await db.enquiries.count_documents({})
    new_enquiries = await db.enquiries.count_documents({"status": "New"})
    customers = await db.customers.count_documents({})
    quotations = await db.quotations.count_documents({})
    invoices = await db.invoices.count_documents({})
    inv_docs = await db.invoices.find({}).to_list(2000)
    total_sales = sum(i.get("total", 0) for i in inv_docs)
    paid = sum(sum(p["amount"] for p in i.get("payments", [])) for i in inv_docs)
    pending = total_sales - paid
    cat_downloads = sum(c.get("downloads", 0) for c in await db.catalogues.find({}).to_list(200))

    months = []
    for k in range(5, -1, -1):
        m = (datetime.now(timezone.utc) - timedelta(days=30 * k)).strftime("%Y-%m")
        months.append(m)
    sales_series = []
    for m in months:
        m_inv = [i for i in inv_docs if i.get("month") == m]
        m_q = await db.quotations.count_documents({"month": m})
        m_e = await db.enquiries.count_documents({"month": m})
        sales_series.append({"month": m, "sales": round(sum(i.get("total", 0) for i in m_inv), 2),
                             "invoices": len(m_inv), "quotations": m_q, "enquiries": m_e})

    recent_enquiries = [clean(e) for e in await db.enquiries.find({}).sort("created_at", -1).limit(5).to_list(5)]
    recent_quotations = [clean(q) for q in await db.quotations.find({}).sort("created_at", -1).limit(5).to_list(5)]
    recent_invoices = [clean(i) for i in await db.invoices.find({}).sort("created_at", -1).limit(5).to_list(5)]
    top_products = [clean(p) for p in await db.products.find({}).sort("views", -1).limit(6).to_list(6)]
    most_enquired = [clean(p) for p in await db.products.find({}).sort("enquiry_count", -1).limit(6).to_list(6)]

    return {"cards": {"total_products": total_products, "active_products": active_products,
                      "categories": categories, "enquiries": enquiries, "new_enquiries": new_enquiries,
                      "customers": customers, "quotations": quotations, "invoices": invoices,
                      "total_sales": round(total_sales, 2), "pending_payments": round(pending, 2),
                      "catalogue_downloads": cat_downloads},
            "monthly": sales_series,
            "recent_enquiries": recent_enquiries, "recent_quotations": recent_quotations,
            "recent_invoices": recent_invoices, "top_products": top_products,
            "most_enquired": most_enquired}


@api.get("/")
async def root():
    return {"message": "KEM Enterprises API"}


app.include_router(api)

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=[o.strip() for o in os.environ.get("CORS_ORIGINS", "http://localhost:3000").split(",") if o.strip()],
    allow_methods=["*"],
    allow_headers=["*"],
)


async def seed_content():
    if await db.faqs.count_documents({}) == 0:
        faqs = [
            ("What products does KEM Enterprises supply?", "We supply premium wire rope fittings, cable grippers, wire ropes, suspension & hanging systems, signage systems, stainless steel & glass hardware, swing/jhula fittings, LED power supplies and fasteners for commercial, architectural and industrial applications."),
            ("Do you provide custom solutions?", "Yes. We engineer bespoke suspension, signage and hardware solutions to your exact specifications. Share your requirement via the enquiry form or WhatsApp and our team will assist."),
            ("What is the minimum order quantity (MOQ)?", "MOQ varies by product and is listed on each product page. For bulk and project requirements, contact us for the best pricing."),
            ("How can I get a price or quotation?", "Prices are shared on enquiry. Click 'Request Quote' or 'Enquire on WhatsApp' on any product, or use the contact form — we typically respond within 24 hours."),
            ("Do you ship across India?", "Yes, we supply to B2B customers across India with reliable lead times. Shipping terms are confirmed at the time of quotation."),
        ]
        await db.faqs.insert_many([
            {"id": str(uuid.uuid4()), "question": q, "answer": a, "order": i, "active": True, "created_at": now_iso()}
            for i, (q, a) in enumerate(faqs)])
    if await db.blog_posts.count_documents({}) == 0:
        posts = [
            ("How to Choose the Right Wire Rope Gripper for Suspension", "A practical guide to selecting cable grippers by load, wire diameter and application.",
             "Cable grippers are the backbone of modern suspension systems. When selecting a gripper, start with the wire rope diameter — most architectural applications use 1mm to 3mm stainless steel rope. Match the gripper's rated load to your fixture weight with a healthy safety margin (we recommend at least 5x). For ceiling signage and lighting, self-locking adjustable grippers make installation fast and clean, while two-side lock grippers give precise levelling. Always use the correct ceiling attachment fitting for your surface — concrete, false ceiling or trunking each need a different top fixing."),
            ("Stainless Steel 304 vs 316: Which Finish Lasts Longer?", "Understand the difference between SS304 and SS316 for indoor and outdoor hardware.",
             "SS304 is the workhorse of interior hardware — excellent corrosion resistance, bright finish and cost-effective for signage, glass studs and railing accessories in dry environments. SS316 adds molybdenum, giving superior resistance to chlorides and marine or humid conditions, making it the right choice for outdoor swings, coastal installations and pool-side railings. For most indoor commercial fit-outs, SS304 is ideal; specify SS316 wherever moisture, salt or chemicals are present."),
            ("Designing Clean Signage Suspension Systems", "Best practices for hanging signage and displays with a premium, minimal look.",
             "The secret to premium signage suspension is invisibility — the hardware should disappear so the sign takes centre stage. Use transparent wire or fine stainless rope with slim top and bottom fittings. Keep tension even across multiple drops, and pre-plan your ceiling attachment points to align with the sign edges. Adjustable grippers let you fine-tune height on-site without tools. For heavier boards, step up to swage stud terminals rated for the full load."),
        ]
        await db.blog_posts.insert_many([
            {"id": str(uuid.uuid4()), "title": t, "slug": slugify(t), "excerpt": e, "content": c,
             "cover_image": "", "author": "KEM Enterprises", "tags": ["guide"], "published": True,
             "created_at": now_iso()} for (t, e, c) in posts])


@app.on_event("startup")
async def startup():
    try:
        init_storage()
        logger.info("Storage initialized")
    except Exception as e:
        logger.error(f"Storage init failed: {e}")
    try:
        await db.users.create_index("email", unique=True)
        await db.products.create_index("slug")
        await db.categories.create_index("slug")
    except Exception as e:
        logger.error(f"Index error: {e}")
    await seed_admin()
    await seed_data()
    await seed_content()
    logger.info("Startup seed complete")


@app.on_event("shutdown")
async def shutdown():
    client.close()
